"""
Génère l'Appendix 2 : caractéristiques des classes CAH 5 clusters
Colonnes : Classe, N communes, Mean Opp, Mean Cho, Mean Vec (Liv), Mean OppChoVec
Deux versions : poids égal et poids Betti
"""
import json, math, sys, pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.stdout.reconfigure(encoding='utf-8')

CAH_JSON  = 'c:/These/oppchovec_visu/stage_ambroise/Code/WEB/cah_5_clusters.json'
DATA_JSON = 'c:/These/oppchovec_visu/stage_ambroise/Code/WEB/data_scores_0_10.json'
OUT       = 'C:/Users/comiti_g/Downloads/appendix2_cah5_classes.xlsx'

# === Charger données ===
with open(CAH_JSON,  encoding='utf-8') as f: cah  = json.load(f)
with open(DATA_JSON, encoding='utf-8') as f: data = json.load(f)

communes = list(data.keys())
alpha, beta = 2.5, 1.5

# === Poids Betti ===
def mean(x): return sum(x)/len(x)
def std(x):  m=mean(x); return math.sqrt(sum((v-m)**2 for v in x)/len(x))
def pearson(a,b):
    n=len(a); ma,mb=mean(a),mean(b)
    num=sum((a[i]-ma)*(b[i]-mb) for i in range(n))
    da=math.sqrt(sum((v-ma)**2 for v in a)); db=math.sqrt(sum((v-mb)**2 for v in b))
    return num/(da*db) if da*db else 0
def cv(x): m=mean(x); return std(x)/m if m else 0
def norm010(d):
    mn,mx=min(d.values()),max(d.values())
    return {c:(d[c]-mn)/(mx-mn)*10 if mx!=mn else 5 for c in communes}

dims = ['Score_Opp','Score_Cho','Score_Vec']
sc01 = {k:[data[c][k] for c in communes] for k in dims}
p1   = [cv(sc01[k]) for k in dims]
corrs= [[pearson(sc01[dims[i]],sc01[dims[j]]) for j in range(3)] for i in range(3)]
p2   = [1/mean([abs(corrs[i][j]) for j in range(3)]) for i in range(3)]
pkRaw= [p1[i]*p2[i] for i in range(3)]; s=sum(pkRaw)
pk_betti=[v/s for v in pkRaw]
pk_egal =[1,1,1]
print(f"Poids Betti : Opp={pk_betti[0]:.4f}  Cho={pk_betti[1]:.4f}  Vec={pk_betti[2]:.4f}")

def calc_occ_norm(pk):
    raw = {c:(1/3)*sum(pk[i]*data[c][dims[i]]**beta for i in range(3))**(alpha/beta) for c in communes}
    return norm010(raw)

occ_egal  = calc_occ_norm(pk_egal)
occ_betti = calc_occ_norm(pk_betti)
sc_opp010 = norm010({c:data[c]['Score_Opp'] for c in communes})
sc_cho010 = norm010({c:data[c]['Score_Cho'] for c in communes})
sc_vec010 = norm010({c:data[c]['Score_Vec'] for c in communes})

# === Construire le tableau ===
clusters_data = cah['clusters']

def build_table(occ_dict, label_occ):
    rows = []
    for cl in sorted(set(v['cluster'] for v in clusters_data.values())):
        membres = [c for c,v in clusters_data.items() if v['cluster'] == cl]
        n = len(membres)
        rows.append({
            'Classe':            cl,
            'N communes':        n,
            'Mean Opp (0-10)':   round(mean([sc_opp010[c] for c in membres]), 2),
            'Mean Cho (0-10)':   round(mean([sc_cho010[c] for c in membres]), 2),
            'Mean Liv (0-10)':   round(mean([sc_vec010[c] for c in membres]), 2),
            f'Mean OppChoLiv ({label_occ})': round(mean([occ_dict[c] for c in membres]), 2),
        })
    return rows

rows_egal  = build_table(occ_egal,  'égal')
rows_betti = build_table(occ_betti, 'Betti')

# === Affichage console ===
print("\n--- Poids ÉGAL ---")
for r in rows_egal:
    print(f"  Classe {r['Classe']} | N={r['N communes']:3d} | "
          f"Opp={r['Mean Opp (0-10)']:.2f}  Cho={r['Mean Cho (0-10)']:.2f}  "
          f"Liv={r['Mean Liv (0-10)']:.2f}  OCC={r['Mean OppChoLiv (égal)']:.2f}")

print("\n--- Poids BETTI ---")
for r in rows_betti:
    print(f"  Classe {r['Classe']} | N={r['N communes']:3d} | "
          f"Opp={r['Mean Opp (0-10)']:.2f}  Cho={r['Mean Cho (0-10)']:.2f}  "
          f"Liv={r['Mean Liv (0-10)']:.2f}  OCC={r['Mean OppChoLiv (Betti)']:.2f}")

# === Écrire Excel ===
df_egal  = pd.DataFrame(rows_egal)
df_betti = pd.DataFrame(rows_betti)

# Ligne totale
def add_total(df, occ_col):
    total = {'Classe': 'Total', 'N communes': df['N communes'].sum()}
    for col in df.columns:
        if col not in ('Classe', 'N communes'):
            total[col] = round((df[col] * df['N communes']).sum() / df['N communes'].sum(), 2)
    return pd.concat([df, pd.DataFrame([total])], ignore_index=True)

df_egal  = add_total(df_egal,  'Mean OppChoLiv (égal)')
df_betti = add_total(df_betti, 'Mean OppChoLiv (Betti)')

with pd.ExcelWriter(OUT, engine='openpyxl') as writer:
    for df, sheet in [(df_egal, 'Égal'), (df_betti, 'Betti')]:
        df.to_excel(writer, sheet_name=sheet, index=False)
        ws = writer.sheets[sheet]

        # Largeurs colonnes
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 14
        for col in range(3, len(df.columns)+1):
            ws.column_dimensions[get_column_letter(col)].width = 18

        # Style en-têtes
        header_fill = PatternFill("solid", fgColor="1F4E79")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', wrap_text=True)

        # Style données + alternance lignes
        fills = [PatternFill("solid", fgColor="EBF3FB"), PatternFill("solid", fgColor="FFFFFF")]
        thin = Side(style='thin', color='BBBBBB')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row)):
            is_total = (row[0].value == 'Total')
            for cell in row:
                cell.alignment = Alignment(horizontal='center')
                cell.border = border
                if is_total:
                    cell.fill = PatternFill("solid", fgColor="D6E4F0")
                    cell.font = Font(bold=True)
                else:
                    cell.fill = fills[i % 2]

print(f"\nOK → {OUT}")
