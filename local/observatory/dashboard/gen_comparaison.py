import json, math, sys, pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.stdout.reconfigure(encoding='utf-8')

with open('c:/These/oppchovec_visu/stage_ambroise/Code/WEB/data_scores_0_10.json', encoding='utf-8') as f:
    data = json.load(f)

communes = list(data.keys())
alpha, beta = 2.5, 1.5

# === Fonctions ===
def mean(x): return sum(x)/len(x)
def std(x): m=mean(x); return math.sqrt(sum((v-m)**2 for v in x)/len(x))
def pearson(a,b):
    n=len(a); ma,mb=mean(a),mean(b)
    num=sum((a[i]-ma)*(b[i]-mb) for i in range(n))
    da=math.sqrt(sum((v-ma)**2 for v in a)); db=math.sqrt(sum((v-mb)**2 for v in b))
    return num/(da*db) if da*db else 0
def cv(x): m=mean(x); return std(x)/m if m else 0
def quantile(x,q):
    s=sorted(x); idx=q*(len(s)-1); lo=int(idx); hi=lo+1
    return s[lo]+(idx-lo)*(s[hi]-s[lo]) if hi<len(s) else s[lo]
def gini(x):
    s = sorted(x); n = len(s); m = mean(s)
    if m == 0: return 0
    return (2*sum((i+1)*s[i] for i in range(n))) / (n*sum(s)) - (n+1)/n

def stats_row(label, vals):
    v=sorted(vals); m=mean(v); s=std(v)
    q1=quantile(v,.25); med=quantile(v,.5); q3=quantile(v,.75)
    g=gini(v)
    return {'Indicateur': label,
            'Nombre': len(v),
            'Moyenne': round(m,6),
            'Ecart-type': round(s,6),
            'Variance': round(s**2,6),
            'Minimum': round(v[0],6),
            'Q1 (25%)': round(q1,6),
            'Mediane': round(med,6),
            'Q3 (75%)': round(q3,6),
            'Maximum': round(v[-1],6),
            'IQR (Q3-Q1)': round(q3-q1,6),
            'Gini': round(g,6)}

def norm010(d):
    mn,mx=min(d.values()),max(d.values())
    return {c:(d[c]-mn)/(mx-mn)*10 if mx!=mn else 5 for c in communes}

# === Poids Betti (sur scores 0-1) ===
dims_keys = ['Score_Opp','Score_Cho','Score_Vec']
sc01 = {k:[data[c][k] for c in communes] for k in dims_keys}
p1=[cv(sc01[k]) for k in dims_keys]
corrs=[[pearson(sc01[dims_keys[i]],sc01[dims_keys[j]]) for j in range(3)] for i in range(3)]
p2=[1/mean([abs(corrs[i][j]) for j in range(3)]) for i in range(3)]
pkRaw=[p1[i]*p2[i] for i in range(3)]; s=sum(pkRaw)
pk_betti=[v/s for v in pkRaw]
pk_egal=[1,1,1]
print(f"Poids Betti: Opp={pk_betti[0]:.4f} Cho={pk_betti[1]:.4f} Vec={pk_betti[2]:.4f}")

def calc_occ(pk):
    raw={}
    for c in communes:
        d=[data[c][k] for k in dims_keys]
        raw[c]=(1/3)*sum(pk[i]*d[i]**beta for i in range(3))**(alpha/beta)
    return raw

raw_egal  = calc_occ(pk_egal)
raw_betti = calc_occ(pk_betti)
occ_egal_010  = norm010(raw_egal)
occ_betti_010 = norm010(raw_betti)

# === Construire le tableau comparatif ===
rows = []

# Séparateur
rows.append({'Indicateur': '── OppChoVec Égal [1,1,1] ──'})
rows.append(stats_row('OppChoVec_0_10  [egal]',  list(occ_egal_010.values())))
rows.append(stats_row('OppChoVec brut  [egal]',  list(raw_egal.values())))

rows.append({'Indicateur': ''})  # ligne vide
rows.append({'Indicateur': f'── OppChoVec Betti (Opp={pk_betti[0]:.3f}, Cho={pk_betti[1]:.3f}, Vec={pk_betti[2]:.3f}) ──'})
rows.append(stats_row('OppChoVec_0_10  [betti]', list(occ_betti_010.values())))
rows.append(stats_row('OppChoVec brut  [betti]', list(raw_betti.values())))

rows.append({'Indicateur': ''})
rows.append({'Indicateur': '── Dimensions (inchangées) ──'})
sc_opp010 = norm010({c:data[c]['Score_Opp'] for c in communes})
sc_cho010 = norm010({c:data[c]['Score_Cho'] for c in communes})
sc_vec010 = norm010({c:data[c]['Score_Vec'] for c in communes})
rows.append(stats_row('Score_Opp_0_10', [sc_opp010[c] for c in communes]))
rows.append(stats_row('Score_Opp',      [data[c]['Score_Opp'] for c in communes]))
rows.append(stats_row('Score_Cho_0_10', [sc_cho010[c] for c in communes]))
rows.append(stats_row('Score_Cho',      [data[c]['Score_Cho'] for c in communes]))
rows.append(stats_row('Score_Vec_0_10', [sc_vec010[c] for c in communes]))
rows.append(stats_row('Score_Vec',      [data[c]['Score_Vec'] for c in communes]))

df = pd.DataFrame(rows).set_index('Indicateur')

out = 'C:/Users/comiti_g/Downloads/comparaison_egal_vs_betti_0_10.xlsx'
with pd.ExcelWriter(out, engine='openpyxl') as writer:
    df.to_excel(writer, sheet_name='Comparaison')

    # Mise en forme
    ws = writer.sheets['Comparaison']
    ws.column_dimensions['A'].width = 46

    # Couleurs
    blue_fill   = PatternFill("solid", fgColor="DDEEFF")
    green_fill  = PatternFill("solid", fgColor="DDFFEE")
    grey_fill   = PatternFill("solid", fgColor="EEEEEE")
    header_font = Font(bold=True)

    for row in ws.iter_rows():
        cell_a = row[0]
        val = str(cell_a.value or '')
        if 'Égal' in val or 'egal' in val.lower():
            for c in row: c.fill = blue_fill; c.font = Font(bold=True)
        elif 'Betti' in val or 'betti' in val.lower():
            for c in row: c.fill = green_fill; c.font = Font(bold=True)
        elif 'Dimensions' in val:
            for c in row: c.fill = grey_fill; c.font = Font(bold=True)

    # Largeur colonnes numériques
    for col in range(2, 13):
        ws.column_dimensions[get_column_letter(col)].width = 14

print(f"OK -> {out}")

# Résumé Gini
print(f"\n=== Gini ===")
for label, vals in [
    ('OppChoVec_0_10 egal ', list(occ_egal_010.values())),
    ('OppChoVec_0_10 betti', list(occ_betti_010.values())),
    ('OppChoVec brut egal ', list(raw_egal.values())),
    ('OppChoVec brut betti', list(raw_betti.values())),
]:
    print(f"  {label}: {gini(vals):.6f}")
